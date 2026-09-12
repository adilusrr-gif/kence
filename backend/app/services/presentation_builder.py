import io
from pathlib import Path
from typing import Optional
from pptx import Presentation as PPTXPresentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR_TYPE, MSO_SHAPE
from app.core.config import get_settings
from app.services.presentation_plan import THEMES

settings = get_settings()


def _hex(color: str) -> RGBColor:
    h = color.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _mix(hex_a: str, hex_b: str, t: float) -> str:
    """Interpolates between two hex colors (t=0 -> a, t=1 -> b)."""
    a, b = hex_a.lstrip('#'), hex_b.lstrip('#')
    ar, ag, ab = int(a[0:2], 16), int(a[2:4], 16), int(a[4:6], 16)
    br, bgn, bb = int(b[0:2], 16), int(b[2:4], 16), int(b[4:6], 16)
    r = round(ar + (br - ar) * t)
    g = round(ag + (bgn - ag) * t)
    bl = round(ab + (bb - ab) * t)
    return f'#{r:02x}{g:02x}{bl:02x}'


def _add_text_box(slide, text: str, left, top, width, height,
                  font_size=18, bold=False, color="#FFFFFF", align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = _hex(color)
    return txBox


def _fill_text_frame(text_frame, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    """lines: list of (text, font_size, bold, color) rendered as stacked paragraphs."""
    text_frame.word_wrap = True
    text_frame.vertical_anchor = anchor
    for i, (text, size, bold, color) in enumerate(lines):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _hex(color)


def _add_card(slide, left, top, width, height, fill_color, radius=0.06):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex(fill_color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    try:
        shape.adjustments[0] = radius
    except (IndexError, AttributeError):
        pass
    return shape


def _fill_solid_bg(slide, color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _hex(color)


def _fill_gradient_bg(slide, color_a: str, color_b: str, angle: float = 45):
    fill = slide.background.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = _hex(color_a)
    stops[1].color.rgb = _hex(color_b)
    try:
        fill.gradient_angle = angle
    except Exception:
        pass


def _add_speaker_notes(slide, points: list, title_text: str):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = "\n".join(points) if points else title_text


def _fetch_logo_bytes(logo_url: Optional[str]) -> Optional[bytes]:
    """Best-effort fetch of an org's branding logo. Never raises — a broken/missing
    logo must not block presentation generation."""
    if not logo_url or not logo_url.startswith(("http://", "https://")):
        return None
    try:
        import httpx
        resp = httpx.get(logo_url, timeout=4.0, follow_redirects=True)
        resp.raise_for_status()
        return resp.content
    except Exception:
        return None


def _add_logo(slide, logo_bytes: bytes, left, top, width):
    try:
        slide.shapes.add_picture(io.BytesIO(logo_bytes), left, top, width=width)
    except Exception:
        pass


def build_presentation(plan: dict, theme_name: str, selected_ids: list, session_id: str,
                        org_id: Optional[int] = None, progress_cb=None) -> str:
    from app.services.chart_service import generate_chart_bytes, generate_infographic_bytes

    theme = THEMES.get(theme_name, THEMES["corporate"])
    bg_color = theme["bg"]
    accent_color = theme["accent"]
    text_color = theme["text"]

    logo_bytes = None
    if org_id is not None:
        try:
            from app.services.branding_service import get_branding
            branding = get_branding(org_id)
            logo_bytes = _fetch_logo_bytes(branding.get("logo_url"))
        except Exception:
            logo_bytes = None

    prs = PPTXPresentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    slides_to_build = [s for s in plan.get("slides", []) if s.get("id") in selected_ids]
    total = len(slides_to_build)
    doc_title = plan.get("title", "")[:50]
    section_counter = 0

    for idx, slide_data in enumerate(slides_to_build, 1):
        stype = slide_data.get("type", "content")
        if progress_cb:
            progress_cb({"status": f"Слайд {idx}/{total}: {slide_data.get('title', '')[:40]}", "slide": idx, "total": total})
        slide = prs.slides.add_slide(blank_layout)

        if stype in ("title", "section_divider"):
            _fill_gradient_bg(slide, bg_color, _mix(bg_color, accent_color, 0.35))
        else:
            _fill_solid_bg(slide, bg_color)

        title_text = slide_data.get("title", "")
        points = slide_data.get("points", [])

        if stype == "title":
            if logo_bytes:
                _add_logo(slide, logo_bytes, Inches(0.6), Inches(0.5), Inches(1.4))
            _add_text_box(slide, title_text,
                          Inches(1.5), Inches(2.6), Inches(10.33), Inches(1.5),
                          font_size=40, bold=True, color=text_color, align=PP_ALIGN.CENTER)
            subtitle = points[0] if points else ""
            if subtitle:
                _add_text_box(slide, subtitle,
                              Inches(2), Inches(4.4), Inches(9.33), Inches(0.8),
                              font_size=18, color=accent_color, align=PP_ALIGN.CENTER)

        elif stype == "section_divider":
            section_counter += 1
            _add_text_box(slide, f"{section_counter:02d}",
                          Inches(1), Inches(2.0), Inches(3.5), Inches(2),
                          font_size=96, bold=True, color=accent_color)
            _add_text_box(slide, title_text,
                          Inches(1), Inches(4.0), Inches(10.5), Inches(1.2),
                          font_size=32, bold=True, color=text_color)
            subtitle = points[0] if points else ""
            if subtitle:
                _add_text_box(slide, subtitle,
                              Inches(1), Inches(4.95), Inches(9.5), Inches(0.7),
                              font_size=15, color=text_color)

        elif stype == "big_number":
            _add_text_box(slide, title_text,
                          Inches(0.8), Inches(0.6), Inches(11.5), Inches(0.7),
                          font_size=20, bold=True, color=accent_color)
            hero = points[0] if points else ""
            _add_text_box(slide, hero,
                          Inches(0.75), Inches(1.9), Inches(11.5), Inches(2.4),
                          font_size=84, bold=True, color=text_color)
            y = Inches(4.6)
            for pt in points[1:4]:
                _add_text_box(slide, f"›  {pt}",
                              Inches(0.9), y, Inches(11), Inches(0.55),
                              font_size=16, color=text_color)
                y += Inches(0.6)

        elif stype == "kpi_row":
            _add_text_box(slide, title_text,
                          Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.7),
                          font_size=26, bold=True, color=text_color)
            data = slide_data.get("data") or {}
            labels = (data.get("labels") or [])[:4]
            values = (data.get("values") or [])[:4]
            unit = data.get("unit", "")
            n = max(len(labels), 1)
            gap = Inches(0.3)
            usable_width = Inches(12.3)
            card_w = int((usable_width - gap * (n - 1)) / n) if n > 1 else usable_width
            left = Inches(0.5)
            for i in range(n):
                card = _add_card(slide, left, Inches(2.2), card_w, Inches(3.1),
                                  _mix(bg_color, accent_color, 0.14))
                value = values[i] if i < len(values) else ""
                value_text = f"{value:g}{(' ' + unit) if unit else ''}" if isinstance(value, (int, float)) else str(value)
                label_text = labels[i] if i < len(labels) else ""
                _fill_text_frame(card.text_frame, [
                    (value_text, 32, True, accent_color),
                    (label_text, 13, False, text_color),
                ])
                left = int(left + card_w + gap)

        elif stype == "two_column":
            _add_text_box(slide, title_text,
                          Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8),
                          font_size=26, bold=True, color=text_color)
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR_TYPE.STRAIGHT, Inches(0.5), Inches(1.3), Inches(12.83), Inches(1.3))
            connector.line.color.rgb = _hex(accent_color)
            connector.line.width = Pt(2)

            col_w = Inches(5.9)
            _add_card(slide, Inches(0.5), Inches(1.55), col_w, Inches(5.55),
                      _mix(bg_color, "#000000", 0.1))
            _add_card(slide, Inches(6.85), Inches(1.55), col_w, Inches(5.55),
                      _mix(bg_color, accent_color, 0.16))

            left_title = slide_data.get("left_title", "")
            right_title = slide_data.get("right_title", "")
            _add_text_box(slide, left_title,
                          Inches(0.8), Inches(1.75), Inches(5.3), Inches(0.6),
                          font_size=18, bold=True, color=text_color)
            _add_text_box(slide, right_title,
                          Inches(7.15), Inches(1.75), Inches(5.3), Inches(0.6),
                          font_size=18, bold=True, color=accent_color)

            y = Inches(2.5)
            for pt in (slide_data.get("left_points") or [])[:5]:
                _add_text_box(slide, f"›  {pt}", Inches(0.8), y, Inches(5.3), Inches(0.6),
                              font_size=14, color=text_color)
                y += Inches(0.6)
            y = Inches(2.5)
            for pt in (slide_data.get("right_points") or [])[:5]:
                _add_text_box(slide, f"›  {pt}", Inches(7.15), y, Inches(5.3), Inches(0.6),
                              font_size=14, color=text_color)
                y += Inches(0.6)

        elif stype == "chart":
            _add_text_box(slide, title_text,
                          Inches(0.5), Inches(0.3), Inches(12), Inches(0.8),
                          font_size=24, bold=True, color=text_color)
            chart_data = slide_data.get("data") or {}
            labels, values = chart_data.get("labels"), chart_data.get("values")
            try:
                if not labels or not values:
                    raise ValueError("No grounded chart data on slide")
                chart_bytes = generate_chart_bytes(
                    chart_type=slide_data.get("chart_type", "bar"),
                    title=title_text,
                    data_hint="",
                    theme=theme,
                    direct_labels=labels,
                    direct_values=values,
                )
                img_stream = io.BytesIO(chart_bytes)
                slide.shapes.add_picture(img_stream, Inches(1.5), Inches(1.4), Inches(10), Inches(5.6))
            except Exception:
                _add_text_box(slide, "[График недоступен]",
                              Inches(1.5), Inches(3), Inches(10), Inches(1),
                              font_size=16, color=accent_color, align=PP_ALIGN.CENTER)

        elif stype == "image":
            _add_text_box(slide, title_text,
                          Inches(0.5), Inches(0.2), Inches(12), Inches(0.8),
                          font_size=24, bold=True, color=text_color)
            try:
                img_bytes = generate_infographic_bytes(title_text, points, theme)
                slide.shapes.add_picture(io.BytesIO(img_bytes),
                                         Inches(0.4), Inches(1.2), Inches(12.5), Inches(5.8))
            except Exception:
                y = Inches(1.6)
                for pt in points[:8]:
                    _add_text_box(slide, f"›  {pt}",
                                  Inches(0.8), y, Inches(11.5), Inches(0.6),
                                  font_size=16, color=text_color)
                    y += Inches(0.65)

        elif stype == "quote":
            _add_text_box(slide, "❝",
                          Inches(1), Inches(1.2), Inches(1), Inches(1),
                          font_size=60, color=accent_color)
            quote_text = points[0] if points else title_text
            _add_text_box(slide, quote_text,
                          Inches(1.5), Inches(2.0), Inches(10), Inches(3),
                          font_size=22, bold=False, color=text_color, align=PP_ALIGN.CENTER)

        else:
            _add_text_box(slide, title_text,
                          Inches(0.5), Inches(0.3), Inches(12), Inches(0.9),
                          font_size=28, bold=True, color=text_color)

            line = slide.shapes.add_connector(
                MSO_CONNECTOR_TYPE.STRAIGHT,
                Inches(0.5), Inches(1.35),
                Inches(12.8), Inches(1.35),
            )
            line.line.color.rgb = _hex(accent_color)
            line.line.width = Pt(2)

            y = Inches(1.6)
            step = Inches(0.65)
            for pt in points[:8]:
                _add_text_box(slide, f"›  {pt}",
                              Inches(0.8), y, Inches(11.5), Inches(0.6),
                              font_size=16, color=text_color)
                y += step

        # Speaker notes for every slide
        _add_speaker_notes(slide, points, title_text)

        # Footer: slide number (bottom-right) + doc title (bottom-left) + logo mark (top-right)
        if stype not in ("title", "section_divider"):
            _add_text_box(slide, f"{idx}/{total}",
                          Inches(12.0), Inches(7.1), Inches(1.2), Inches(0.35),
                          font_size=11, color=accent_color, align=PP_ALIGN.RIGHT)
            if doc_title:
                _add_text_box(slide, doc_title,
                              Inches(0.4), Inches(7.1), Inches(8), Inches(0.35),
                              font_size=10, color=text_color)
            if logo_bytes:
                _add_logo(slide, logo_bytes, Inches(12.5), Inches(0.25), Inches(0.55))

    # Closing slide — added automatically, not part of the user's slide count/footer.
    if slides_to_build:
        if progress_cb:
            progress_cb({"status": "Финальный слайд…"})
        closing = prs.slides.add_slide(blank_layout)
        _fill_gradient_bg(closing, bg_color, _mix(bg_color, accent_color, 0.35))
        if logo_bytes:
            _add_logo(closing, logo_bytes, Inches(5.87), Inches(2.0), Inches(1.6))
        _add_text_box(closing, "Спасибо за внимание",
                      Inches(1.5), Inches(3.6), Inches(10.33), Inches(1.1),
                      font_size=34, bold=True, color=text_color, align=PP_ALIGN.CENTER)
        if doc_title:
            _add_text_box(closing, doc_title,
                          Inches(2), Inches(4.7), Inches(9.33), Inches(0.7),
                          font_size=15, color=accent_color, align=PP_ALIGN.CENTER)

    if progress_cb:
        progress_cb({"status": "Сохраняю файл…"})
    out_path = Path(settings.UPLOAD_DIR) / session_id / "presentation_v2.pptx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return str(out_path)
