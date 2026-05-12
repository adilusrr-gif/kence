from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path
from typing import Dict, List
import json

class PresentationGenerator:
    def __init__(self):
        self.template_colors = {
            "primary": RGBColor(0x1A, 0x5F, 0x7A),      # Тёмно-бирюзовый
            "secondary": RGBColor(0x57, 0xC5, 0xB6),    # Бирюзовый
            "accent": RGBColor(0xFF, 0x6B, 0x6B),       # Коралловый
            "text": RGBColor(0x2C, 0x3E, 0x50),         # Тёмно-синий
            "light": RGBColor(0xF8, 0xF9, 0xFA)         # Белый/светлый
        }

    def generate(self, structure: dict, session_id: str) -> str:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        slides_data = structure.get("slides", [])
        title = structure.get("title", "Презентация")

        # Титульный слайд
        self._add_title_slide(prs, title, structure.get("subtitle", ""))

        # Контентные слайды
        for slide_data in slides_data[1:] if len(slides_data) > 1 else slides_data:
            self._add_content_slide(prs, slide_data)

        # Сохранение
        output_path = Path(f"./uploads/{session_id}/presentation.pptx")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))

        return str(output_path)

    def _add_title_slide(self, prs: Presentation, title: str, subtitle: str):
        blank_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(blank_layout)

        # Фон
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.template_colors["primary"]

        # Заголовок
        title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(54)
        p.font.bold = True
        p.font.color.rgb = self.template_colors["light"]
        p.alignment = PP_ALIGN.CENTER

        # Подзаголовок
        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.333), Inches(1))
            tf = sub_box.text_frame
            p = tf.paragraphs[0]
            p.text = subtitle
            p.font.size = Pt(24)
            p.font.color.rgb = self.template_colors["secondary"]
            p.alignment = PP_ALIGN.CENTER

    def _add_content_slide(self, prs: Presentation, slide_data: dict):
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)

        # Фон
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.template_colors["light"]

        # Полоса сверху
        bar = slide.shapes.add_shape(
            1, Inches(0), Inches(0), Inches(13.333), Inches(0.15)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.template_colors["primary"]
        bar.line.fill.background()

        # Заголовок слайда
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(1))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = slide_data.get("title", "")
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = self.template_colors["primary"]

        # Пункты
        points = slide_data.get("points", [])
        content_box = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(11.333), Inches(5))
        tf = content_box.text_frame
        tf.word_wrap = True

        for i, point in enumerate(points):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()

            p.text = f"• {point}"
            p.font.size = Pt(20)
            p.font.color.rgb = self.template_colors["text"]
            p.space_after = Pt(16)
            p.level = 0

pptx_generator = PresentationGenerator()
