"""Document Editor Agent — edits full document content, builds charts, exports DOCX."""
import io
import re
import logging
from pathlib import Path
from app.services.agents.base_agent import AgentState, emit_step
from app.core.session import session_manager
from app.core.config import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()

_CHART_RE = re.compile(
    r'\[CHART\s+type=["\']?(\w+)["\']?\s+title=["\']?([^"\']+?)["\']?\s+data_hint=["\']?([^"\']+?)["\']?\]',
    re.IGNORECASE,
)

_EDIT_PROMPT = """Ты — профессиональный редактор документов. Ниже дан документ и инструкции по редактированию.

ИНСТРУКЦИИ:
{instructions}

ДОКУМЕНТ:
{document}

ПРАВИЛА:
- Выполни все запрошенные изменения (исправь текст, улучши таблицы, добавь разделы)
- Для вставки графика используй ТОЛЬКО формат: [CHART type="bar" title="Заголовок" data_hint="краткое описание данных"]
  Допустимые типы: bar, line, pie
- Сохраняй markdown-форматирование: # заголовки, | таблицы |, **жирный**, - списки
- Верни ТОЛЬКО исправленный документ без каких-либо объяснений или комментариев

ИСПРАВЛЕННЫЙ ДОКУМЕНТ:"""


_MAX_INSTRUCTIONS_LEN = 2000

_INJECTION_PATTERNS = re.compile(
    r'(ignore\s+(previous|all|prior)|system\s*:|assistant\s*:|<\|.*?\|>|'
    r'forget\s+(everything|instructions)|new\s+instructions\s*:|'
    r'you\s+are\s+now|disregard\s+(all|previous)|'
    r'ignore\s+all\s+prior)',
    re.IGNORECASE,
)


def _sanitize_instructions(text: str) -> str:
    """Trim and strip prompt-injection patterns from user instructions."""
    text = text.strip()
    if len(text) > _MAX_INSTRUCTIONS_LEN:
        text = text[:_MAX_INSTRUCTIONS_LEN]
    text = _INJECTION_PATTERNS.sub("[filtered]", text)
    return text


async def run(state: AgentState, llm_service) -> AgentState:
    task_id = state["task_id"]
    session_id = state.get("session_id")
    raw_instructions = (
        state.get("instructions") or state.get("question")
        or "Улучши структуру и читаемость документа"
    )
    instructions = _sanitize_instructions(raw_instructions)

    # 1. Load full document content
    emit_step(task_id, "load_document", "running", "Загрузка документа...")
    session = session_manager.get_session(session_id)
    if not session:
        return {**state, "error": f"Session not found: {session_id}"}

    full_text = session.get("markdown_text", "")
    if not full_text:
        return {**state, "error": "Документ не содержит текста. Загрузите документ перед редактированием."}

    emit_step(task_id, "load_document", "done", f"Загружено {len(full_text)} символов")

    # 2. LLM edits the full document
    emit_step(task_id, "edit_document", "running", "Редактирование с помощью LLM...")
    prompt = _EDIT_PROMPT.format(document=full_text, instructions=instructions)
    try:
        edited_markdown = await llm_service.agenerate(prompt)
    except Exception as e:
        return {**state, "error": f"LLM error: {e}"}

    emit_step(task_id, "edit_document", "done",
              f"Готово. Изменение размера: {len(edited_markdown) - len(full_text):+d} символов")

    # 3. Render [CHART ...] directives
    chart_matches = list(_CHART_RE.finditer(edited_markdown))
    charts: dict[str, bytes] = {}

    if chart_matches:
        emit_step(task_id, "generate_charts", "running",
                  f"Генерация {len(chart_matches)} графиков...")
        from app.services.chart_service import generate_chart_bytes
        theme = {"bg": "#FFFFFF", "accent": "#2563EB", "text": "#1E293B"}
        for m in chart_matches:
            directive = m.group(0)
            if directive in charts:
                continue
            try:
                charts[directive] = generate_chart_bytes(
                    m.group(1), m.group(2), m.group(3), theme, llm_service
                )
            except Exception as e:
                logger.warning("Chart generation failed for '%s': %s", m.group(2), e)
        emit_step(task_id, "generate_charts", "done",
                  f"Построено {len(charts)} графиков")
    else:
        emit_step(task_id, "generate_charts", "done", "Графики не запрошены")

    # 4. Export to DOCX with tables + embedded chart images
    emit_step(task_id, "export_docx", "running", "Создание DOCX...")
    try:
        output_path = _build_docx(session_id, edited_markdown, charts)
        emit_step(task_id, "export_docx", "done", f"Файл: {output_path.name}")
    except Exception as e:
        logger.error("DOCX export failed: %s", e)
        return {**state, "error": f"Export failed: {e}"}

    download_url = f"/api/agents/tasks/{task_id}/download"

    return {
        **state,
        "result": {
            "answer": (
                f"Документ отредактирован. "
                f"Изменение: {len(edited_markdown) - len(full_text):+d} символов. "
                f"Графиков: {len(charts)}."
            ),
            "edited_markdown": edited_markdown,
            "download_path": str(output_path),
            "download_url": download_url,
            "charts_count": len(charts),
        },
    }


def _build_docx(session_id: str, markdown: str, charts: dict[str, bytes]) -> Path:
    """Build a DOCX from edited markdown with embedded charts and rendered tables."""
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    output_path = Path(settings.UPLOAD_DIR) / session_id / "edited_document.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    lines = markdown.split("\n")
    table_buffer: list[str] = []
    i = 0

    while i < len(lines):
        line = lines[i].rstrip()

        # Flush accumulated table when non-table line arrives
        if table_buffer and not line.startswith("|"):
            _render_table(doc, table_buffer)
            table_buffer = []

        # Table row
        if line.startswith("|"):
            table_buffer.append(line)
            i += 1
            continue

        # Chart directive
        cm = _CHART_RE.match(line)
        if cm:
            directive = cm.group(0)
            if directive in charts:
                try:
                    doc.add_picture(io.BytesIO(charts[directive]), width=Inches(5.5))
                    doc.add_paragraph("")
                except Exception:
                    doc.add_paragraph(f"[График: {cm.group(2)}]")
            else:
                doc.add_paragraph(f"[График: {cm.group(2)}]")
            i += 1
            continue

        # Standard markdown
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r'^\d+\. ', line):
            doc.add_paragraph(re.sub(r'^\d+\. ', '', line), style="List Number")
        elif line in ("", "---"):
            doc.add_paragraph("")
        else:
            clean = re.sub(r'\*{1,2}([^*]+)\*{1,2}', r'\1', line)
            clean = re.sub(r'`([^`]+)`', r'\1', clean)
            if clean.strip():
                doc.add_paragraph(clean)

        i += 1

    if table_buffer:
        _render_table(doc, table_buffer)

    doc.save(str(output_path))
    return output_path


def _render_table(doc, table_lines: list[str]):
    """Convert markdown table lines into a python-docx table."""
    data_rows = [l for l in table_lines if not re.match(r'^\|[\s\-:|]+\|$', l)]
    if not data_rows:
        return
    rows = [[c.strip() for c in l.strip('|').split('|')] for l in data_rows]
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    try:
        table.style = 'Table Grid'
    except Exception:
        pass
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            if c_idx < n_cols:
                cell = table.cell(r_idx, c_idx)
                cell.text = cell_text
                if r_idx == 0:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
    doc.add_paragraph("")
